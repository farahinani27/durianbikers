#!/usr/bin/env python3
"""
Read one tab of the Qbee stock masterlist (xlsx export) into CSV or JSON.

Handles the parts that are identical for every tab and easy to get wrong:
  - skips hidden rows (archived/discontinued -- ~2,326 of them on HELMETS)
  - casts float-formatted stock ("1.0" -> 1) instead of regex-matching digits
  - stops at a banner row such as "OUT OF STOCK RANGE"
  - carries hierarchical BRAND/SERIES cells downward (blank = same as above)
  - keeps the original sheet row number on every record, so anything odd can be
    traced back to a cell

Read-only. It never writes to the workbook.

Usage:
    python3 read_tab.py MASTERLIST.xlsx --list
    python3 read_tab.py MASTERLIST.xlsx HELMETS --header-row 2 --carry 0,1 \
        --stop-at 'OUT OF STOCK RANGE' --max-col 11

Note: cell notes/comments are NOT exported reliably by Google Sheets -- on the
HELMETS tab 68% are mis-anchored to the wrong row. This script deliberately does
not read them. See ../references/helmets.md.
"""

import argparse
import csv
import json
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl not installed. Run: python3 -m pip install openpyxl")


def as_stock(value):
    """Stock cells arrive as floats ('1.0'). Return an int, or None if not numeric."""
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_tab(path, tab, header_row, stop_at, carry_cols, include_hidden, max_col):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if tab not in wb.sheetnames:
        sys.exit("Tab %r not found. Available: %s" % (tab, ", ".join(wb.sheetnames)))
    ws = wb[tab]

    hidden = {i for i, rd in ws.row_dimensions.items() if rd.hidden}

    headers = []
    if header_row:
        raw = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        headers = [clean(c) for c in raw]
    if max_col:
        headers = headers[:max_col]

    rows = []
    carried = {}
    skipped_hidden = 0
    stopped_at_row = None

    start = (header_row + 1) if header_row else 1
    for idx, raw in enumerate(ws.iter_rows(min_row=start, values_only=True), start=start):
        cells = [clean(c) for c in raw]
        if max_col:
            cells = cells[:max_col]

        if stop_at and any(stop_at.upper() in c.upper() for c in cells if c):
            stopped_at_row = idx
            break

        if idx in hidden:
            skipped_hidden += 1
            if not include_hidden:
                continue

        for col in carry_cols:
            if col < len(cells):
                if cells[col]:
                    carried[col] = cells[col]
                elif col in carried:
                    cells[col] = carried[col]

        if not any(cells):
            continue

        rows.append({"row": idx, "hidden": idx in hidden, "cells": cells})

    return {
        "tab": tab,
        "header_row": header_row,
        "headers": headers,
        "rows": rows,
        "skipped_hidden": skipped_hidden,
        "stopped_at_row": stopped_at_row,
        "total_hidden_in_sheet": len(hidden),
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("tab", nargs="?")
    ap.add_argument("--list", action="store_true", help="list tab names and exit")
    ap.add_argument("--header-row", type=int, default=1,
                    help="1-indexed header row (HELMETS uses 2). 0 = no header")
    ap.add_argument("--stop-at", default=None,
                    help="stop at the first row containing this text, e.g. 'OUT OF STOCK RANGE'")
    ap.add_argument("--carry", default="",
                    help="comma-separated 0-indexed columns where blank means 'same as above' "
                         "(HELMETS BRAND/SERIES = 0,1)")
    ap.add_argument("--include-hidden", action="store_true",
                    help="keep hidden rows (they are archived/discontinued -- normally skip)")
    ap.add_argument("--max-col", type=int, default=0, help="truncate to the first N columns")
    ap.add_argument("--format", choices=["csv", "json"], default="csv")
    args = ap.parse_args()

    if args.list:
        wb = openpyxl.load_workbook(args.workbook, read_only=True)
        print("\n".join(wb.sheetnames))
        return

    if not args.tab:
        ap.error("tab is required (or use --list)")

    carry_cols = [int(c) for c in args.carry.split(",") if c.strip() != ""]
    result = read_tab(args.workbook, args.tab, args.header_row, args.stop_at,
                      carry_cols, args.include_hidden, args.max_col)

    if args.format == "json":
        json.dump(result, sys.stdout, indent=2)
        sys.stdout.write("\n")
    else:
        out = csv.writer(sys.stdout)
        out.writerow(["_row"] + (result["headers"] or []))
        for r in result["rows"]:
            out.writerow([r["row"]] + r["cells"])

    sys.stderr.write(
        "tab=%s rows=%d hidden_skipped=%d hidden_total=%d stopped_at_row=%s\n"
        % (result["tab"], len(result["rows"]), result["skipped_hidden"],
           result["total_hidden_in_sheet"], result["stopped_at_row"])
    )


if __name__ == "__main__":
    main()
